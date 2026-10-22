"""Import a Scrivener multi-document markdown dump into a project. (2.3)

Flow of a dump import:

    read file ─ split on ===DOCUMENT=== ─ blank-line header fields ─ RawDocument list
                                                              │
                              book folders = parent titles that are not documents
                                                              │
                    import_book: folder → stack walk → add_document / write_content /
                    set_metadata / update_document(synopsis)

The dump format is flat: every block carries Title, Type, Parent (a
title, not an id), Position and a body after ===TEXT===. The blocks are
in depth-first binder order, which is what makes parent-by-title safe
even though titles repeat across books: walking the list with a stack of
open ancestors resolves each Parent to the nearest still-open ancestor
with that title — verified against the author's full corpus with zero
orphans before this code was written.

Only two kinds appear in the author's dumps (chapter, scene); anything
else lands as a scene, and folders are never in the dump at all — book
folders and intermediate containers are rebuilt here, not imported.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from html import escape
from pathlib import Path

import docx

from .project import Project

# Every key the author's Scrivener compile emits. Continuation chunks in
# a multi-line field (a Synopsis written as story beats, for instance)
# are recognised by NOT being one of these keys, so the whitelist is
# load-bearing rather than documentation.
HEADER_KEYS = (
    "Type", "ID", "Hierarchy", "Position", "ParentPosition", "Parent",
    "Level1", "Level2", "Title", "PoV", "Location", "Date", "Quote",
    "Source", "Status", "Label", "Keywords", "Synopsis",
)

# Headers that describe the dump itself, not the document: used to build
# the tree, then dropped. Hierarchy is Scrivener's compile placeholder
# (always the same junk token in the author's dumps) and is dropped too.
STRUCTURAL_KEYS = {
    "Type", "ID", "Hierarchy", "Position", "ParentPosition", "Parent",
    "Level1", "Level2", "Title", "Synopsis",
}

# Scrivener writes these sentinels when a field was never filled in.
# Storing them would bury the real values under 300 "No Status" rows.
UNSET_VALUES = {"No Status", "No Label", "None", ""}

_FIELD_LINE = re.compile(r"^([A-Za-z][A-Za-z0-9]*):[ \t]?(.*)$")


@dataclass
class RawDocument:
    """One ===DOCUMENT=== block, parsed but not yet placed."""

    title: str
    kind: str                       # "chapter", "scene" — the dump's Type
    parent: str                     # parent's Title as declared ("" = top)
    position: int                   # sibling order inside the parent
    synopsis: str
    body: str                       # markdown-ish prose after ===TEXT===
    metadata: dict[str, str] = field(default_factory=dict)
    value_types: dict[str, str] = field(default_factory=dict)


def parse_dump(path: Path) -> list[RawDocument]:
    """Parse a multi-document markdown dump into RawDocuments, in dump order.

    Dump order is depth-first binder order and is the contract every
    later step relies on — reordering the result would break parent
    resolution.
    """
    text = Path(path).read_text(encoding="utf-8")
    documents: list[RawDocument] = []
    for block in text.split("===DOCUMENT===")[1:]:
        if "===TEXT===" not in block:
            continue
        header_part, body = block.split("===TEXT===", 1)
        fields = _parse_header(header_part)
        documents.append(RawDocument(
            title=fields.get("Title", "").strip() or "Untitled",
            kind=fields.get("Type", "Scene").strip().lower(),
            parent=fields.get("Parent", "").strip(),
            position=_as_int(fields.get("Position", "0")),
            synopsis=fields.get("Synopsis", "").strip(),
            body=body.strip("\n"),
            metadata=_kept_metadata(fields),
            value_types=_value_types(fields),
        ))
    return documents


def _parse_header(header_part: str) -> dict[str, str]:
    """Blank-line-separated Key: value chunks → fields.

    A chunk whose first line is not a known key continues the previous
    field (multi-line Synopsis values arrive as bare continuation
    chunks), so an unknown chunk never invents a field.
    """
    fields: dict[str, str] = {}
    current: str | None = None
    for chunk in re.split(r"\n\s*\n", header_part):
        lines = [line for line in chunk.splitlines() if line.strip()]
        if not lines:
            continue
        match = _FIELD_LINE.match(lines[0].strip())
        if match and match.group(1) in HEADER_KEYS:
            current = match.group(1)
            fields[current] = match.group(2).strip()
            continuation = [line.strip() for line in lines[1:]]
        else:
            continuation = [line.strip() for line in lines]
        if current is not None and continuation:
            joined = fields[current] + "\n" + "\n".join(continuation)
            fields[current] = joined
    return fields


def _kept_metadata(fields: dict[str, str]) -> dict[str, str]:
    """The document-metadata keys worth storing: everything non-structural
    that the writer actually filled in."""
    kept: dict[str, str] = {}
    for key, value in fields.items():
        if key in STRUCTURAL_KEYS:
            continue
        cleaned = value.strip()
        if cleaned in UNSET_VALUES:
            continue
        kept[key] = cleaned
    return kept


def _value_types(fields: dict[str, str]) -> dict[str, str]:
    """Types the dump's Date column as a date so it arrives as a picker.

    The author asked for "dates changed to dates directly": an imported
    date should open in the panel as a date widget showing the stored
    day, not as text to retype. Every other kept key stays plain text.
    """
    return {
        key: "date"
        for key, value in fields.items()
        if key not in STRUCTURAL_KEYS
        and key.strip().lower() == "date"
        and value.strip() not in UNSET_VALUES
    }


def _as_int(text: str) -> int:
    try:
        return int(text.strip())
    except (TypeError, ValueError):
        return 0


def shelf_works(raw_documents: list[RawDocument]) -> dict[str, list[str]]:
    """Containers that are a shelf of works rather than a book.

    A top-level parent (its title matches no document — only such
    parents can be books) with at least two direct children, every one
    of them a chapter, is Scrivener's generic "Manuscript" arrangement:
    separate works stored as chapter-level documents. Real books in the
    author's corpus always have a book-level scene among their children
    (a title page), so all-chapters never misfires on them. Returns
    shelf title -> work titles, in binder order.
    """
    titles = {document.title for document in raw_documents}
    children: dict[str, list[RawDocument]] = {}
    for document in raw_documents:
        if document.parent and document.parent not in titles:
            children.setdefault(document.parent, []).append(document)
    return {
        parent: [child.title for child in kids]
        for parent, kids in children.items()
        if len(kids) >= 2 and all(child.kind == "chapter" for child in kids)
    }


def dump_books(raw_documents: list[RawDocument]) -> list[str]:
    """The books a dump contains, in binder order.

    A book is a parent title that is not itself a document — the dump
    never includes folder blocks, so the books only exist as the parents
    their chapters hang off — except for shelves: each work a shelf
    holds becomes its own book, since that is what the author means by
    a book ("the third book should be Perfectly Flawed…").
    """
    shelves = shelf_works(raw_documents)
    titles = {document.title for document in raw_documents}
    books: list[str] = []
    for document in raw_documents:
        parent = document.parent
        if not parent or parent in titles:
            continue
        for book in shelves.get(parent, [parent]):
            if book not in books:
                books.append(book)
    return books


def markdown_to_html(text: str) -> str:
    """Minimal markdown → the rich-text HTML Lekhak stores.

    Deliberately tiny: paragraphs split on blank lines, ``*em*`` and
    ``**strong**``, leading ``#`` heading lines (the dump's title-page
    blocks use them) rendered bold. Everything is escaped first and
    anything fancier stays visible as text — this is an import of prose,
    not a pandoc replacement; the writer polishes in the editor.
    """
    paragraphs: list[str] = []
    for chunk in re.split(r"\n\s*\n", text.strip()):
        lines = [line.rstrip() for line in chunk.splitlines() if line.strip()]
        if not lines:
            continue
        joined = "\n".join(lines)
        heading = re.match(r"^#{1,6}\s+(.*?)\s*#*\s*$", joined)
        if heading:
            paragraphs.append(f"<p><b>{_inline(heading.group(1))}</b></p>")
            continue
        escaped = _inline(joined).replace("\n", "<br>")
        paragraphs.append(f"<p>{escaped}</p>")
    return "\n".join(paragraphs)


def _inline(text: str) -> str:
    """Escape HTML, then turn the two markdown inline markers into tags."""
    escaped = (text.replace("&", "&amp;").replace("<", "&lt;")
               .replace(">", "&gt;"))
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"\*(.+?)\*", r"<i>\1</i>", escaped)
    return escaped


def markdown_from_html(html: str) -> str:
    """Inverse of markdown_to_html's subset (2.3): rich-text HTML back
    to clean markdown prose.

    Paragraphs become blank-line-separated blocks, ``<br>`` back to a
    newline, ``<b>``/``<i>`` back to ``**``/``*``; every other tag is
    dropped and the four usual entities unescaped. Accepts whatever Qt
    stored — callers pass content through ``accepted_html`` first, so
    tracking marks and style spans are already gone — and aims only at
    the shapes this app's own converters put in. A literal ``*`` left
    in prose can misparse on a later import; the writer polishes.
    """
    html = re.sub(r"<style\b[^>]*>.*?</style>", " ", html, flags=re.I | re.S)
    html = re.sub(r"<br\s*/?>", "\n", html, flags=re.I)
    html = re.sub(r"</p\s*>", "\n\n", html, flags=re.I)
    html = re.sub(r"<(?:b|strong)>", "**", html, flags=re.I)
    html = re.sub(r"</(?:b|strong)>", "**", html, flags=re.I)
    html = re.sub(r"<(?:i|em)>", "*", html, flags=re.I)
    html = re.sub(r"</(?:i|em)>", "*", html, flags=re.I)
    html = re.sub(r"<[^>]+>", "", html)
    html = (html.replace("&nbsp;", " ").replace("&amp;", "&")
            .replace("&lt;", "<").replace("&gt;", ">"))
    # Stored paragraphs are joined by a newline as well as a closing
    # tag, which would otherwise leave three between blocks.
    html = re.sub(r"\n{3,}", "\n\n", html)
    return html.strip()


def import_book(project: Project, raw_documents: list[RawDocument],
                book_title: str, parent_id: str | None = None) -> str:
    """Create one book from a parsed dump; returns the book folder's id.

    Walks the raw list in dump order with an ancestor stack, so scenes
    land under their own book's chapter even when another book has a
    chapter with the same title. Documents whose parent chain never
    reaches ``book_title`` belong to another book and are skipped.

    A promoted shelf work (see shelf_works) differs in one way: the
    work's own chapter block — whose parent is the shelf, not the book —
    becomes the book's first child, and its scenes hang under that
    chapter exactly as they did in Scrivener, so no structure and no
    prose is lost in the promotion.
    """
    shelf_names = set(shelf_works(raw_documents))
    folder = project.add_document(
        project.unique_title(book_title, parent_id), parent_id, "folder")
    stack: list[tuple[str, str]] = []          # (title, document_id)
    work_chapter: str | None = None
    for raw in raw_documents:
        while stack and stack[-1][0] != raw.parent:
            stack.pop()
        if raw.parent == book_title:
            child_parent = work_chapter or folder.id
        elif raw.title == book_title and raw.parent in shelf_names:
            child_parent = folder.id
        elif stack:
            child_parent = stack[-1][1]
        else:
            stack = []      # not in this book; a later block may re-enter
            continue
        kind = raw.kind if raw.kind in ("chapter", "scene") else "scene"
        # Containers (chapters) must not land on a name already there: a
        # same-titled second chapter would make every name-based match
        # ambiguous. Scenes keep their titles — repeated scene titles are
        # ordinary, and an import has no business renaming prose.
        raw_title = (project.unique_title(raw.title, child_parent)
                     if kind == "chapter" else raw.title)
        document = project.add_document(raw_title, child_parent, kind)
        if raw.synopsis:
            project.update_document(document.id, synopsis=raw.synopsis)
        if raw.metadata:
            project.set_metadata(document.id, raw.metadata, raw.value_types)
        project.write_content(document.id, markdown_to_html(raw.body))
        stack.append((raw.title, document.id))
        if raw.title == book_title and raw.parent in shelf_names:
            work_chapter = document.id
    return folder.id


def import_all_books(project: Project, raw_documents: list[RawDocument]
                     ) -> tuple[str | None, list[str]]:
    """Import every book in a dump; returns (wrapper id or None, book ids).

    This is the normal path: the author hands over one file containing
    their whole library and expects Lekhak to do the splitting, so the
    caller never has to know what books the dump contains.

    When the dump holds exactly one shelf (the author's "Manuscript"),
    the shelf becomes a wrapper folder and every book lands inside it —
    folder blocks are absent from the dump, so the shelf is the only
    trace of where the books lived, and the author's reference
    structure confirms Marginally Human and Terminally-Immortal sit
    inside Manuscript beside the works. With no shelf, or more than
    one, books import at the top level instead: nothing in the dump
    can say which shelf an outside book belongs to.
    """
    books = dump_books(raw_documents)
    shelves = shelf_works(raw_documents)
    wrapper_id: str | None = None
    if len(shelves) == 1:
        wrapper_id = project.add_document(
            project.unique_title(next(iter(shelves)), None), None, "folder"
        ).id
    return wrapper_id, [
        import_book(project, raw_documents, book, wrapper_id)
        for book in books
    ]


# --- plain markdown import (2.3) ---------------------------------------------
# The author's own proposal: the writer names the divider strings in the
# import dialog — no heading convention is imposed — and the file splits
# on them, pandoc-free, everything between dividers becoming prose.


@dataclass
class MarkdownSplit:
    """One document's worth of a split markdown file."""

    kind: str                       # "chapter" | "scene"
    title: str
    body: str


def _divider_kind(line: str, chapter_divider: str,
                  scene_divider: str) -> tuple[str | None, str]:
    """Does this line open a new document? Returns (kind, title).

    The scene divider is tested first so the more specific string wins
    when both would match (the '#' / '##' defaults). A divider must sit
    at the start of a stripped line and be followed by end-of-line or
    whitespace, so '##nope' is prose rather than a scene called 'nope'.
    """
    stripped = line.strip()
    for kind, divider in (("scene", scene_divider),
                          ("chapter", chapter_divider)):
        if not divider or not stripped.startswith(divider):
            continue
        rest = stripped[len(divider):]
        if rest and not rest[0].isspace():
            continue
        return kind, rest.strip()
    return None, ""


def split_markdown(text: str, *, chapter_divider: str = "#",
                   scene_divider: str = "##",
                   opening_title: str = "Opening") -> list[MarkdownSplit]:
    """Split a plain markdown file on the writer's own divider strings.

    Prose before the first divider becomes a scene of its own (titled by
    the caller — the import names it after the book), so nothing is
    silently dropped when a file opens with a paragraph instead of a
    heading. Blank lines inside a body survive, ready for
    markdown_to_html's paragraph split.
    """
    chunks: list[MarkdownSplit] = []
    kind: str | None = None
    title = opening_title
    body: list[str] = []
    for line in text.splitlines():
        found_kind, found_title = _divider_kind(
            line, chapter_divider, scene_divider)
        if found_kind:
            if kind is not None:
                chunks.append(MarkdownSplit(kind, title, "\n".join(body)))
            kind, title, body = found_kind, found_title or "Untitled", []
        elif kind is not None:
            body.append(line)
        else:
            kind, title, body = "scene", opening_title, [line]
    if kind is not None:
        chunks.append(MarkdownSplit(kind, title, "\n".join(body)))
    return chunks


def import_markdown_file(project: Project, path: Path, *,
                         book_title: str = "", chapter_divider: str = "#",
                         scene_divider: str = "##") -> str:
    """Import one plain markdown file as a book; returns the folder id.

    The file becomes a folder named after itself (or the caller's
    title); chapter dividers become chapters, scene dividers become
    scenes under the current chapter — or directly under the book when
    none has started. Bodies go through the same minimal markdown
    conversion the dump import uses.
    """
    path = Path(path)
    title = book_title.strip() or path.stem
    text = path.read_text(encoding="utf-8")
    chunks = split_markdown(text, chapter_divider=chapter_divider,
                            scene_divider=scene_divider,
                            opening_title=title)
    folder = project.add_document(project.unique_title(title, None), None, "folder")
    chapter_id: str | None = None
    for chunk in chunks:
        kind = chunk.kind if chunk.kind in ("chapter", "scene") else "scene"
        document = project.add_document(chunk.title, chapter_id or folder.id,
                                        kind)
        if kind == "chapter":
            chapter_id = document.id
        project.write_content(document.id, markdown_to_html(chunk.body))
    return folder.id


# --- Word / OpenDocument manuscript import (BUILD-NOW 7) ----------------
# An ODT reaches this code by first becoming a DOCX (odt.convert), so
# both formats share one heading map and one import path.

OPENING_TITLE = "Opening"


@dataclass
class DocxChunk:
    """One heading-delimited span: a chapter, a scene, or the opening run."""

    kind: str          # "chapter" (heading 1) or "scene" (deeper heading / prose)
    title: str
    html: str


@dataclass
class DocxManuscript:
    """A parsed Word document: its Title property and its chunks."""

    title: str
    chunks: list[DocxChunk]


def _heading_level(paragraph) -> int:
    """1 for a chapter heading, 2+ for a scene heading, 0 for prose.

    Matched against both the style id and its display name, so a
    localized Word document (name "Überschrift 1", id "Heading1") and a
    LibreOffice round trip land in the same place.
    """
    style = getattr(paragraph, "style", None)
    if style is None:
        return 0
    for token in (getattr(style, "style_id", "") or "",
                  getattr(style, "name", "") or ""):
        match = re.match(r"(?i)^heading[\s_-]?(\d)", token.strip())
        if match:
            return int(match.group(1))
    return 0


def _paragraph_html(paragraph) -> str:
    """One paragraph as the editor's HTML; empty when there is no text.

    Bold, italic and underline ride inside their runs — the three the
    format bar offers, nested in that order. Text that lives outside
    the runs (fields, hyperlink text in some writers) falls back to the
    paragraph's own text, so nothing silently disappears.
    """
    if not paragraph.text.strip():
        return ""
    pieces = []
    for run in paragraph.runs:
        text = escape(run.text).replace("\n", "<br/>")
        if not text:
            continue
        if run.bold:
            text = f"<strong>{text}</strong>"
        if run.italic:
            text = f"<em>{text}</em>"
        if run.underline:
            text = f"<u>{text}</u>"
        pieces.append(text)
    if not pieces:
        return f"<p>{escape(paragraph.text).replace(chr(10), '<br/>')}</p>"
    return "<p>" + "".join(pieces) + "</p>"


def parse_docx(path: Path) -> DocxManuscript:
    """Read a Word document into heading-driven chunks.

    Heading 1 opens a chapter; any deeper heading opens a scene under
    the chapter in force (or under the book before one starts); every
    other paragraph accumulates as prose into whichever container is
    open, starting an "Opening" scene when prose comes before any
    heading. A heading's text becomes its container's binder title and
    never its content — the rule the markdown splitter follows for
    divider lines.
    """
    document = docx.Document(str(path))
    title = (document.core_properties.title or "").strip()
    chunks: list[DocxChunk] = []
    current: DocxChunk | None = None
    for paragraph in document.paragraphs:
        level = _heading_level(paragraph)
        if level:
            current = DocxChunk(
                "chapter" if level == 1 else "scene",
                paragraph.text.strip() or "Untitled", "")
            chunks.append(current)
            continue
        html = _paragraph_html(paragraph)
        if not html:
            continue
        if current is None:
            current = DocxChunk("scene", OPENING_TITLE, "")
            chunks.append(current)
        current.html += html
    return DocxManuscript(title, chunks)


def import_docx(project: Project, path: Path, *, title: str = "") -> str:
    """Import a Word (or converted ODT) manuscript as a book; folder id.

    The folder takes the document's Title property, else the caller's
    title, else the filename. Chapters sit flat under the book — the
    markdown loop stacks each chapter under the previous one, which is
    not the binder shape anyone asked for — and scenes land under the
    current chapter, or under the book while no chapter has started.
    Containers go through unique_title, so a second import is a numbered
    copy rather than an overwrite; the identical-content warning ahead
    of the real run belongs to the floor in app.py.
    """
    path = Path(path)
    parsed = parse_docx(path)
    folder_title = title.strip() or parsed.title or path.stem
    folder = project.add_document(
        project.unique_title(folder_title, None), None, "folder")
    chapter_id: str | None = None
    for chunk in parsed.chunks:
        if chunk.kind == "chapter":
            document = project.add_document(
                project.unique_title(chunk.title, folder.id), folder.id,
                "chapter")
            chapter_id = document.id
        else:
            document = project.add_document(
                chunk.title, chapter_id or folder.id, "scene")
        project.write_content(document.id, chunk.html)
    return folder.id
