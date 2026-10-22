"""BUILD-NOW 7: DOCX/ODT manuscript import through the existing floor.

Acceptance from docs/BUILD-NOW.md: a Word document imports as
book → chapter → scene using its heading styles; importing it twice
warns "identical"; nothing existing is touched.

The ODT half never parses ODT: ``odt.convert`` drives the same UNO
worker seam PDF export uses (faked here — the suite never launches
LibreOffice, per tests/test_pdf.py), and everything afterwards is
``import_docx`` on the converted file. LibreOffice itself was exercised
by hand against real .odt files while this code was written.
"""

import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import docx  # noqa: E402
from PySide6.QtCore import QSettings  # noqa: E402
from PySide6.QtWidgets import (  # noqa: E402
    QApplication, QFileDialog, QMessageBox,
)

from lekhak import odt  # noqa: E402
from lekhak.app import MainWindow  # noqa: E402
from lekhak.importer import (  # noqa: E402
    OPENING_TITLE, import_docx, parse_docx,
)
from lekhak.project import Project  # noqa: E402

_application = QApplication.instance() or QApplication([])
_application.setOrganizationName("Lekhak")
_application.setApplicationName("Lekhak")
_settings_dir = tempfile.mkdtemp(prefix="lekhak-docx-settings-")
QSettings.setPath(QSettings.NativeFormat, QSettings.UserScope, _settings_dir)


def _make_docx(path: Path, *, title: str | None = "My Novel") -> Path:
    document = docx.Document()
    if title:
        document.core_properties.title = title
    document.add_heading("Chapter One", 1)
    paragraph = document.add_paragraph()
    paragraph.add_run("First ")
    bold = paragraph.add_run("prose")
    bold.bold = True
    document.add_heading("Second Scene", 2)
    document.add_paragraph("Scene prose.")
    document.save(str(path))
    return path


def _documents_by_title(project):
    return {d.title: d for d in project.documents()}


class ImportDocxTests(unittest.TestCase):
    """The parser and the placement, without a window anywhere."""

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.file = _make_docx(self.root / "manuscript.docx")
        self.project = Project.create(self.root / "Novel", "Novel")

    def tearDown(self):
        self.project.close()
        self.temporary.cleanup()

    def _parents(self):
        documents = self.project.documents()
        titles = {d.id: d.title for d in documents}
        return [(d.kind, d.title, titles.get(d.parent_id))
                for d in documents]

    def test_heading_styles_build_book_chapter_scene(self):
        folder_id = import_docx(self.project, self.file)

        documents = _documents_by_title(self.project)
        folder = documents["My Novel"]
        chapter = documents["Chapter One"]
        scene = documents["Second Scene"]
        self.assertEqual(folder.id, folder_id)
        self.assertEqual(folder.kind, "folder")
        self.assertEqual(chapter.kind, "chapter")
        self.assertEqual(chapter.parent_id, folder_id)
        self.assertEqual(scene.kind, "scene")
        self.assertEqual(scene.parent_id, chapter.id)
        self.assertIn("<strong>prose</strong>",
                      self.project.read_content(chapter.id))
        self.assertIn("<p>Scene prose.</p>",
                      self.project.read_content(scene.id))

    def test_a_heading_is_a_title_never_content(self):
        import_docx(self.project, self.file)

        documents = _documents_by_title(self.project)
        self.assertNotIn("Chapter One",
                         self.project.read_content(documents["Chapter One"].id))
        self.assertNotIn("Second Scene",
                         self.project.read_content(documents["Second Scene"].id))

    def test_chapters_sit_flat_under_the_book(self):
        """Unlike the markdown loop, which stacks each chapter under the
        previous one — a second import must read as a binder, not a
        chain."""
        path = self.root / "three.docx"
        document = docx.Document()
        for number in ("One", "Two", "Three"):
            document.add_heading(f"Chapter {number}", 1)
            document.add_paragraph("body")
        document.save(str(path))

        folder_id = import_docx(self.project, path)

        for title in ("Chapter One", "Chapter Two", "Chapter Three"):
            self.assertEqual(_documents_by_title(self.project)[title].parent_id,
                             folder_id, title)

    def test_prose_before_any_heading_opens_the_book(self):
        path = self.root / "cold.docx"
        document = docx.Document()
        document.add_paragraph("A cold open.")
        document.add_heading("Chapter One", 1)
        document.add_paragraph("Body.")
        document.save(str(path))

        folder_id = import_docx(self.project, path)

        opening = _documents_by_title(self.project)[OPENING_TITLE]
        self.assertEqual(opening.kind, "scene")
        self.assertEqual(opening.parent_id, folder_id)
        self.assertIn("A cold open.", self.project.read_content(opening.id))

    def test_a_folderless_file_falls_back_to_its_filename(self):
        path = self.root / "untitled draft.docx"
        document = docx.Document()          # no Title property at all
        document.add_paragraph("Just prose.")
        document.save(str(path))

        folder_id = import_docx(self.project, path)

        documents = _documents_by_title(self.project)
        self.assertEqual(documents["untitled draft"].id, folder_id)
        self.assertEqual(documents[OPENING_TITLE].parent_id, folder_id)

    def test_parse_reads_the_title_property(self):
        self.assertEqual(parse_docx(self.file).title, "My Novel")


class ImportThroughTheAppTests(unittest.TestCase):
    """The File-menu handler over the floor: warn, number, never touch."""

    def setUp(self):
        QSettings().remove("session/lastProject")
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.file = _make_docx(self.root / "manuscript.docx")
        self.project = Project.create(self.root / "Novel", "Novel")
        self.window = MainWindow()
        self.window._switch_project(self.project)

    def tearDown(self):
        self.window.project = None
        self.project.close()
        self.temporary.cleanup()

    def _import_once(self):
        with patch.object(QFileDialog, "getOpenFileName",
                          return_value=(str(self.file), "")):
            self.window.import_docx_odt()

    def _folder_count(self):
        return sum(1 for d in self.project.documents() if d.kind == "folder")

    def test_the_dialog_lands_the_book_and_says_so(self):
        self._import_once()

        self.assertIn("My Novel", _documents_by_title(self.project))
        self.assertIn("Imported manuscript.docx",
                      self.window.statusBar().currentMessage())

    def test_importing_twice_warns_identical_and_adds_nothing(self):
        self._import_once()
        before = self._folder_count()

        with patch.object(QFileDialog, "getOpenFileName",
                          return_value=(str(self.file), "")), \
                patch.object(QMessageBox, "warning",
                             return_value=QMessageBox.No) as warning:
            self.window.import_docx_odt()

        self.assertEqual(self._folder_count(), before)
        self.assertIn("same content", warning.call_args.args[2])

    def test_insisting_keeps_both_copies_and_touches_neither_prose(self):
        self._import_once()
        original = _documents_by_title(self.project)["Chapter One"]
        original_html = self.project.read_content(original.id)

        with patch.object(QFileDialog, "getOpenFileName",
                          return_value=(str(self.file), "")), \
                patch.object(QMessageBox, "warning",
                             return_value=QMessageBox.Yes):
            self.window.import_docx_odt()

        documents = _documents_by_title(self.project)
        self.assertIn("My Novel (2)", documents)
        self.assertIn("Chapter One", documents)
        self.assertEqual(self.project.read_content(original.id), original_html)

    def test_cancelling_the_file_dialog_changes_nothing(self):
        before = self._folder_count()

        with patch.object(QFileDialog, "getOpenFileName",
                          return_value=("", "")), \
                patch.object(QMessageBox, "warning") as warning:
            self.window.import_docx_odt()

        self.assertEqual(self._folder_count(), before)
        warning.assert_not_called()

    def test_without_a_project_the_dialog_is_never_opened(self):
        self.window.project = None

        with patch.object(QFileDialog, "getOpenFileName") as chooser:
            self.window.import_docx_odt()

        chooser.assert_not_called()


class OdtConvertTests(unittest.TestCase):
    """odt.convert standing on pdf.py's seam, with LibreOffice faked."""

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "book.odt"
        self.source.write_bytes(b"fake-odt")

    def tearDown(self):
        self.temporary.cleanup()

    def _run(self, side_effect):
        with patch.object(odt, "find_converter",
                          return_value=["/usr/bin/soffice"]), \
                patch.object(odt, "_find_uno_interpreter",
                             return_value="/usr/bin/python3"), \
                patch.object(odt, "_run_uno", side_effect=side_effect):
            destination = self.root / "out.docx"
            odt.convert(self.source, destination)
        return destination

    def test_the_result_lands_and_the_worker_gets_its_own_script(self):
        seen = {}

        def fake_run_uno(interpreter, source, destination, soffice,
                         workdir, script=None):
            seen.update(interpreter=interpreter, source=source,
                        script=script, workdir=workdir)
            Path(destination).write_bytes(b"converted-docx")
            return subprocess.CompletedProcess([], 0, "", "")

        destination = self._run(fake_run_uno)

        self.assertEqual(destination.read_bytes(), b"converted-docx")
        self.assertEqual(seen["script"], odt.UNO_SCRIPT)
        self.assertEqual(seen["script"].name, "uno_odt.py")
        self.assertEqual(seen["source"], self.source)
        self.assertEqual(seen["interpreter"], "/usr/bin/python3")
        leftovers = [p.name for p in self.root.iterdir()
                     if p.name.startswith(".lekhak-odt")]
        self.assertEqual(leftovers, [])

    def test_a_failed_worker_quotes_libreoffice(self):
        def failing(*args, **kwargs):
            return subprocess.CompletedProcess([], 1, "", "boom")

        with patch.object(odt, "find_converter",
                          return_value=["/usr/bin/soffice"]), \
                patch.object(odt, "_find_uno_interpreter",
                             return_value="/usr/bin/python3"), \
                patch.object(odt, "_run_uno", side_effect=failing):
            with self.assertRaises(odt.ConversionError) as caught:
                odt.convert(self.source, self.root / "out.docx")

        self.assertIn("failed to convert", str(caught.exception))
        self.assertIn("boom", str(caught.exception))

    def test_missing_converter_explains_the_install(self):
        with patch.object(odt, "find_converter", return_value=None):
            with self.assertRaises(odt.ConversionError) as caught:
                odt.convert(self.source, self.root / "out.docx")

        self.assertIn("LibreOffice on PATH", str(caught.exception))

    def test_missing_uno_explains_python3_uno(self):
        with patch.object(odt, "find_converter",
                          return_value=["/usr/bin/soffice"]), \
                patch.object(odt, "_find_uno_interpreter", return_value=None):
            with self.assertRaises(odt.ConversionError) as caught:
                odt.convert(self.source, self.root / "out.docx")

        self.assertIn("python3-uno", str(caught.exception))

    def test_a_timeout_gives_up_readably(self):
        def hang(*args, **kwargs):
            raise subprocess.TimeoutExpired("uno_odt.py", 300)

        with patch.object(odt, "find_converter",
                          return_value=["/usr/bin/soffice"]), \
                patch.object(odt, "_find_uno_interpreter",
                             return_value="/usr/bin/python3"), \
                patch.object(odt, "_run_uno", side_effect=hang):
            with self.assertRaises(odt.ConversionError) as caught:
                odt.convert(self.source, self.root / "out.docx")

        self.assertIn("did not finish", str(caught.exception))

    def test_the_worker_script_exists_and_states_its_usage(self):
        from lekhak import uno_odt
        worker = Path(uno_odt.__file__)

        result = subprocess.run([sys.executable, str(worker)],
                                capture_output=True, text=True, timeout=60)

        self.assertIn("usage: uno_odt.py", result.stderr)


class OdtThroughTheAppTests(unittest.TestCase):
    """The .odt branch of the handler: convert once, then everything
    the DOCX path does."""

    def setUp(self):
        QSettings().remove("session/lastProject")
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.file = self.root / "book.odt"
        self.file.write_bytes(b"fake-odt")
        self.project = Project.create(self.root / "Novel", "Novel")
        self.window = MainWindow()
        self.window._switch_project(self.project)

    def tearDown(self):
        self.window.project = None
        self.project.close()
        self.temporary.cleanup()

    def _fake_convert(self, source, destination):
        # What LibreOffice would hand back: a real DOCX to parse.
        document = docx.Document()
        document.add_heading("Converted Chapter", 1)
        document.add_paragraph("Body after the round trip.")
        document.save(str(destination))

    def test_converted_once_even_though_the_floor_runs_twice(self):
        with patch.object(QFileDialog, "getOpenFileName",
                          return_value=(str(self.file), "")), \
                patch.object(odt, "convert",
                             side_effect=self._fake_convert) as convert:
            self.window.import_docx_odt()

        self.assertEqual(convert.call_count, 1)
        documents = _documents_by_title(self.project)
        self.assertIn("Converted Chapter", documents)
        self.assertEqual(documents["Converted Chapter"].kind, "chapter")

    def test_an_unconvertible_odt_reports_and_lands_nothing(self):
        before = len(self.project.documents())

        def refuse(*args, **kwargs):
            raise odt.ConversionError("LibreOffice said no")

        with patch.object(QFileDialog, "getOpenFileName",
                          return_value=(str(self.file), "")), \
                patch.object(odt, "convert", side_effect=refuse), \
                patch.object(QMessageBox, "critical") as critical:
            self.window.import_docx_odt()

        critical.assert_called_once()
        self.assertIn("LibreOffice said no", critical.call_args.args[2])
        self.assertEqual(len(self.project.documents()), before)


if __name__ == "__main__":
    unittest.main()
